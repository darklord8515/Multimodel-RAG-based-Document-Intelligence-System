"""Lightweight file -> RAG-Anything ``content_list`` converters.

The upstream RAG-Anything pipeline relies on MinerU/Docling for parsing, which
pull in multi-GB ML dependencies. For the web app we convert the most common
formats ourselves and hand the result to ``RAGAnything.insert_content_list``,
so the full multimodal pipeline (text + tables + images) still runs.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

TEXT_EXTS = {".txt", ".md", ".markdown", ".log"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED_EXTS = TEXT_EXTS | IMAGE_EXTS | {".pdf", ".docx", ".csv"}


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [[(c or "").replace("\n", " ").strip() for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(norm[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(lines)


def _parse_pdf(path: Path) -> list[dict[str, Any]]:
    from pypdf import PdfReader

    items = []
    for idx, page in enumerate(PdfReader(str(path)).pages):
        text = (page.extract_text() or "").strip()
        if text:
            items.append({"type": "text", "text": text, "page_idx": idx})
    if not items:
        raise ValueError(
            "No extractable text found. Scanned PDFs need OCR, which the lightweight parser doesn't do."
        )
    return items


def _parse_docx(path: Path) -> list[dict[str, Any]]:
    from docx import Document

    doc = Document(str(path))
    items: list[dict[str, Any]] = []
    para_buf: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            para_buf.append(p.text.strip())
    if para_buf:
        items.append({"type": "text", "text": "\n\n".join(para_buf), "page_idx": 0})
    for table in doc.tables:
        rows = [[c.text for c in row.cells] for row in table.rows]
        body = _md_table(rows)
        if body:
            items.append(
                {"type": "table", "table_body": body, "table_caption": [], "table_footnote": [], "page_idx": 0}
            )
    if not items:
        raise ValueError("The document appears to be empty.")
    return items


def _parse_csv(path: Path) -> list[dict[str, Any]]:
    rows = list(csv.reader(io.StringIO(_read_text(path))))
    if not rows:
        raise ValueError("The CSV file is empty.")
    return [
        {
            "type": "table",
            "table_body": _md_table(rows),
            "table_caption": [path.name],
            "table_footnote": [],
            "page_idx": 0,
        }
    ]


def to_content_list(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        text = _read_text(path).strip()
        if not text:
            raise ValueError("The file is empty.")
        return [{"type": "text", "text": text, "page_idx": 0}]
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        return _parse_docx(path)
    if ext == ".csv":
        return _parse_csv(path)
    if ext in IMAGE_EXTS:
        return [
            {
                "type": "image",
                "img_path": str(path.resolve()),
                "image_caption": [path.name],
                "image_footnote": [],
                "page_idx": 0,
            }
        ]
    raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
